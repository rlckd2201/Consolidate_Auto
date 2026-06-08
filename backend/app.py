from __future__ import annotations

import json
import os
import re
import shutil
import socket
import urllib.error
import urllib.request
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Literal

from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from backend.reply_import import load_latest, run_reply_import, save_run

try:
    import pymysql
except ImportError:  # pragma: no cover - dependency is optional until HR lookup is used.
    pymysql = None

try:
    import openpyxl
except ImportError:  # pragma: no cover - dependency is optional until report reconciliation is used.
    openpyxl = None


APP_VERSION = "0.3.2"
ROOT = Path(__file__).resolve().parent
STATIC_DIR = ROOT / "static"
DATA_DIR = ROOT / "data"
STATE_PATH = DATA_DIR / "app_state.json"
GENERATED_DIR = ROOT / "generated"

DEFAULT_TEMPLATE_DIR = ROOT.parent.parent / "보고자료"
DEFAULT_OUTPUT_ROOT = Path(os.environ.get("APPDATA", str(ROOT))) / "Consolidate_Auto"
TEMPLATE_DIR = Path(os.environ.get("CONSOLIDATE_TEMPLATE_DIR", str(DEFAULT_TEMPLATE_DIR)))
OUTPUT_ROOT = Path(os.environ.get("CONSOLIDATE_OUTPUT_ROOT", str(DEFAULT_OUTPUT_ROOT)))
REPLY_DIR = Path(os.environ.get("CONSOLIDATE_REPLY_DIR", str(ROOT.parent.parent / "회신자료")))
IMPORT_RUN_DIR = Path(os.environ.get("CONSOLIDATE_IMPORT_RUN_DIR", str(OUTPUT_ROOT / "import_runs")))
REFERENCE_WORKBOOK_ENV = os.environ.get("CONSOLIDATE_REFERENCE_WORKBOOK", "").strip()
HR_DB_NAME = os.environ.get("HR_DB_NAME", "ksystem_yundong")
GEMINI_MODEL = os.environ.get("GEMINI_MODEL", "gemini-2.5-flash")
GEMINI_API_BASE = os.environ.get("GEMINI_API_BASE", "https://generativelanguage.googleapis.com/v1beta")
APP_FEATURES = [
    "hr_lookup",
    "template_copy",
    "approval_preview",
    "gemini_evidence_triage",
    "reply_import_pipeline",
    "import_review_workspace",
    "import_candidate_quality_gate",
    "ppt_pages_1_3_focus",
    "import_source_evidence_matching",
    "ppt_report_workspace",
    "report_reconciliation_workspace",
    "import_latest_null_guard",
]
REPORT_CATEGORIES = [
    "유지보수",
    "생산 돌대응",
    "납품용기/제품관리",
    "개발/공정개선",
    "고객사 납품대응",
    "행정업무",
    "설비/시설 청소",
    "생산지원",
    "기타",
]
REPORT_FACTORY_ORDER = [
    "D1공장",
    "D2공장",
    "D3공장",
    "P1공장",
    "P2공장",
    "P3공장",
    "P4공장",
    "일강1공장",
    "일강2공장",
    "더원공장",
    "제이엠공장",
    "경영지원",
    "영업",
    "구매",
    "생기",
]
REPORT_BASELINES = {
    "2026-05-W5": {
        "total_headcount": 188,
        "weekend_headcount": 95,
        "reference_headcount": 93,
        "row_count": 150,
        "by_date": {
            "2026-05-30": 71,
            "2026-05-31": 24,
            "2026-06-03": 93,
        },
        "by_category": {
            "유지보수": 35,
            "생산 돌대응": 30,
            "납품용기/제품관리": 21,
            "개발/공정개선": 20,
            "고객사 납품대응": 19,
            "행정업무": 17,
            "설비/시설 청소": 12,
            "생산지원": 10,
            "기타": 24,
        },
        "by_factory": {
            "D1공장": 1,
            "D2공장": 2,
            "D3공장": 13,
            "P2공장": 37,
            "P3공장": 17,
            "P4공장": 4,
            "일강1공장": 2,
            "일강2공장": 52,
            "더원공장": 4,
            "경영지원": 4,
            "영업": 19,
            "구매": 1,
            "생기": 32,
        },
    }
}
BLANK_IMPORT_VALUES = {"", "N/A", "NA", "None", "none", "null", "-", "미기재"}

HR_ENTITY_RULES = {
    "daeseung": {"label": "대승", "binum": "1", "factories": ["D1공장", "D2공장", "D3공장"]},
    "daeseung_precision": {"label": "대승정밀", "binum": "2", "factories": ["P1공장", "P2공장", "P3공장", "P4공장"]},
    "ilgang": {"label": "일강", "binum": "3", "factories": ["일강 1공장", "일강 2공장"]},
    "theone": {"label": "더원", "binum": "4", "factories": ["더원공장"]},
    "jm": {"label": "제이엠", "binum": "4", "factories": ["제이엠공장"]},
}

PRODUCTION_MARKERS = ("생산", "조립", "가공", "주조", "단조", "B/CAP", "C/ROD", "JOINT", "SPIDER", "후처리", "M.P.I", "금형", "제관")
INDIRECT_MARKERS = ("생산관리", "생관", "생기", "생산기술", "제조기술", "품질", "보전", "물류", "자재", "구매", "영업", "인사", "총무", "재정", "전산", "원가", "공구", "개발", "지원", "안전", "환경", "출하")


class Period(BaseModel):
    id: str
    label: str
    start_date: str
    end_date: str
    weekend_dates: list[str]
    reference_dates: list[str]
    status: str = "open"


class LegalEntity(BaseModel):
    entity_code: str
    display_name: str
    factory_label: str
    default_overtime_form_label: str
    business_no: str = ""
    active: bool = True
    note: str = ""


class OvertimeRow(BaseModel):
    id: str
    date: str
    company: str
    source_factory: str
    target_factory: str = ""
    job_group: Literal["직접직", "간접직", "관리직"]
    team: str
    name: str = ""
    position: str = ""
    headcount: int = Field(ge=0)
    hours: float = Field(ge=0)
    category1: str
    category2: str = ""
    detail: str
    reason: str = ""
    confidence: Literal["high", "medium", "low"] = "high"
    review_status: Literal["mapped", "needs_review", "reviewed"] = "mapped"
    source: str = "seed"


class Submission(BaseModel):
    id: str
    period_id: str
    submitter: str
    entity_code: str
    factory: str
    team: str
    status: Literal["draft", "submitted_for_admin_review", "returned", "admin_reviewed", "locked_for_report"] = "draft"
    groupware_draft_status: Literal["not_filled", "draft_filled", "draft_fill_failed", "adapter_pending"] = "not_filled"
    rows: list[OvertimeRow]
    attachments: list[str] = []
    updated_at: str


class SubmissionCreate(BaseModel):
    submitter: str = "담당자"
    entity_code: str = "daeseung"
    factory: str = "D3공장"
    team: str = "생산"
    rows: list[OvertimeRow]


class State(BaseModel):
    periods: list[Period]
    legal_entities: list[LegalEntity]
    submissions: list[Submission]


class HREmployeeCandidate(BaseModel):
    emp_id: str
    name: str
    entity_code: str
    entity_name: str
    factory: str
    db_factory: str
    department: str
    department_code: str
    org_name: str = ""
    org_code: str = ""
    position: str = ""
    duty: str = ""
    job_family: str = ""
    direct_marker: str = ""
    pay_type: str = ""
    job_group: Literal["직접직", "간접직", "관리직"]
    classification_reason: str
    confidence: Literal["high", "medium", "low"] = "high"


class GeminiEvidenceItem(BaseModel):
    source_id: str = ""
    source_file: str = ""
    source_sheet_or_page: str = ""
    source_location: str = ""
    text: str


class GeminiEvidenceAnalyzeRequest(BaseModel):
    task: str = "overtime_reply_evidence_triage"
    period_label: str = ""
    context: dict[str, str] = Field(default_factory=dict)
    evidence: list[GeminiEvidenceItem]


class ReplyImportRunRequest(BaseModel):
    reply_dir: str = ""
    period_label: str = "2026년 5월 5주차"
    allow_gemini: bool = False
    max_ai_sources: int = Field(default=24, ge=0, le=120)
    ai_batch_size: int = Field(default=3, ge=1, le=8)


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _seed_state() -> State:
    period = Period(
        id="2026-05-W5",
        label="2026년 5월 5주차",
        start_date="2026-05-30",
        end_date="2026-06-03",
        weekend_dates=["2026-05-30", "2026-05-31"],
        reference_dates=["2026-06-03"],
    )
    entities = [
        LegalEntity(entity_code="daeseung", display_name="대승", factory_label="D1/D2/D3공장", default_overtime_form_label="대승 - 특근계획서"),
        LegalEntity(entity_code="daeseung_precision", display_name="대승정밀", factory_label="P1/P2/P3/P4공장", default_overtime_form_label="대승정밀 - 특근계획서"),
        LegalEntity(entity_code="theone", display_name="더원", factory_label="더원공장", default_overtime_form_label="더원 - 특근계획서", business_no="421-86-02723"),
        LegalEntity(entity_code="ilgang", display_name="일강", factory_label="일강1/2공장", default_overtime_form_label="일강 - 특근계획서"),
        LegalEntity(entity_code="jm", display_name="제이엠", factory_label="제이엠공장", default_overtime_form_label="제이엠 - 특근계획서", business_no="125-81-54876", note="현재 특근 발생 가능성 낮음"),
    ]
    rows = [
        OvertimeRow(id="r-001", date="2026-05-30", company="대승", source_factory="D2공장", job_group="관리직", team="보전", name="송인섭", position="마스터", headcount=1, hours=4, category1="유지보수", category2="설비유지보수", detail="승용컷팅 콜렛척 가이드 교체"),
        OvertimeRow(id="r-002", date="2026-05-30", company="대승", source_factory="D2공장", job_group="관리직", team="물류", name="박광덕", position="매니저", headcount=1, hours=4, category1="납품대응", category2="고객사 납품", detail="고객사 납품대응 및 창고정리"),
        OvertimeRow(id="r-003", date="2026-05-30", company="대승", source_factory="D3공장", job_group="간접직", team="보전", headcount=2, hours=8, category1="유지보수", category2="설비유지보수", detail="PHB 세척기, RO수 Lay-OUT 변경 및 정상화"),
        OvertimeRow(id="r-004", date="2026-06-03", company="대승정밀", source_factory="P3공장", job_group="관리직", team="품질", name="김박동", position="책임", headcount=1, hours=8, category1="납품용기/제품관리", category2="제품조사", detail="개선누 TAP BURR 밀림 조사"),
        OvertimeRow(id="r-005", date="2026-06-03", company="대승정밀", source_factory="P3공장", job_group="관리직", team="생관", name="고두형", position="매니저", headcount=1, hours=8, category1="기타", category2="재물조사", detail="P4공장 재물실사", confidence="medium"),
        OvertimeRow(id="r-006", date="2026-05-31", company="일강", source_factory="일강2공장", job_group="관리직", team="가공생산", name="최태호", position="마스터", headcount=1, hours=8, category1="생산지원", category2="가공생산", detail="가공라인 지원"),
    ]
    submission = Submission(
        id="sub-001",
        period_id=period.id,
        submitter="샘플 담당자",
        entity_code="daeseung",
        factory="D3공장",
        team="생산",
        rows=rows,
        attachments=["D2공장 특근계획서.pdf", "D3공장 특근 상세.xlsx"],
        updated_at=_now(),
    )
    return State(periods=[period], legal_entities=entities, submissions=[submission])


def load_state() -> State:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not STATE_PATH.exists():
        state = _seed_state()
        save_state(state)
        return state
    state = State.model_validate(json.loads(STATE_PATH.read_text(encoding="utf-8")))
    apply_master_updates(state)
    return state


def save_state(state: State) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    STATE_PATH.write_text(state.model_dump_json(indent=2), encoding="utf-8")


def apply_master_updates(state: State) -> None:
    updates = {
        "theone": {"display_name": "더원", "factory_label": "더원공장", "default_overtime_form_label": "더원 - 특근계획서", "business_no": "421-86-02723"},
        "jm": {"display_name": "제이엠", "factory_label": "제이엠공장", "default_overtime_form_label": "제이엠 - 특근계획서", "business_no": "125-81-54876", "note": "현재 특근 발생 가능성 낮음"},
    }
    existing = {entity.entity_code: entity for entity in state.legal_entities}
    for code, values in updates.items():
        entity = existing.get(code)
        if entity:
            for key, value in values.items():
                setattr(entity, key, value)
        else:
            state.legal_entities.append(LegalEntity(entity_code=code, active=True, **values))


def all_rows(state: State) -> list[OvertimeRow]:
    rows: list[OvertimeRow] = []
    for submission in state.submissions:
        rows.extend(submission.rows)
    return rows


def summarize(rows: list[OvertimeRow], period: Period | None = None) -> dict:
    total = sum(row.headcount for row in rows)
    weekend_dates = set(period.weekend_dates if period else [])
    weekend_total = sum(row.headcount for row in rows if row.date in weekend_dates)
    reference_total = total - weekend_total
    by_category: dict[str, int] = {}
    by_factory: dict[str, int] = {}
    by_date: dict[str, int] = {}
    unresolved = 0
    for row in rows:
        by_category[row.category1] = by_category.get(row.category1, 0) + row.headcount
        by_factory[row.source_factory] = by_factory.get(row.source_factory, 0) + row.headcount
        by_date[row.date] = by_date.get(row.date, 0) + row.headcount
        if row.review_status == "needs_review" or row.confidence == "low":
            unresolved += 1
    return {
        "total_headcount": total,
        "weekend_headcount": weekend_total,
        "reference_headcount": reference_total,
        "by_category": by_category,
        "by_factory": by_factory,
        "by_date": by_date,
        "unresolved_rows": unresolved,
    }


def split_counts(rows: list[OvertimeRow]) -> dict[str, int]:
    counts: dict[str, int] = {"직접직": 0, "간접직": 0, "관리직": 0}
    for row in rows:
        counts[row.job_group] = counts.get(row.job_group, 0) + row.headcount
    return counts


def hr_db_configured() -> bool:
    required = ("HR_DB_HOST", "HR_DB_USER", "HR_DB_PASSWORD")
    return all(os.environ.get(key) for key in required)


def gemini_configured() -> bool:
    return gemini_key_status()["ok"]


def gemini_key_status() -> dict:
    api_key = os.environ.get("GEMINI_API_KEY", "")
    if not api_key:
        return {"ok": False, "status": "missing"}
    if "<" in api_key or ">" in api_key or "키" in api_key or "secret" in api_key.lower():
        return {"ok": False, "status": "placeholder_value"}
    try:
        api_key.encode("ascii")
    except UnicodeEncodeError:
        return {"ok": False, "status": "non_ascii_value"}
    if len(api_key.strip()) < 20:
        return {"ok": False, "status": "too_short"}
    return {"ok": True, "status": "configured"}


def gemini_analysis_schema() -> dict:
    candidate_row_schema = {
        "type": "object",
        "properties": {
            "date": {"type": "string"},
            "source_factory": {"type": "string"},
            "target_factory": {"type": "string"},
            "team": {"type": "string"},
            "name": {"type": "string"},
            "job_group": {"type": "string"},
            "headcount": {"type": "string"},
            "hours": {"type": "string"},
            "detail": {"type": "string"},
            "category_guess": {"type": "string"},
            "confidence": {"type": "string"},
            "evidence": {"type": "string"},
            "needs_review_reason": {"type": "string"},
        },
        "required": ["date", "source_factory", "target_factory", "team", "name", "job_group", "headcount", "hours", "detail", "category_guess", "confidence", "evidence", "needs_review_reason"],
    }
    source_schema = {
        "type": "object",
        "properties": {
            "source_id": {"type": "string"},
            "source_type": {"type": "string"},
            "role_guess": {"type": "string"},
            "confidence": {"type": "string"},
            "evidence": {"type": "string"},
            "warning": {"type": "string"},
        },
        "required": ["source_id", "source_type", "role_guess", "confidence", "evidence", "warning"],
    }
    return {
        "type": "object",
        "properties": {
            "summary": {"type": "string"},
            "source_assessment": {"type": "array", "items": source_schema},
            "candidate_rows": {"type": "array", "items": candidate_row_schema},
            "conflicts": {"type": "array", "items": {"type": "string"}},
            "questions_for_human": {"type": "array", "items": {"type": "string"}},
            "abstentions": {"type": "array", "items": {"type": "string"}},
        },
        "required": ["summary", "source_assessment", "candidate_rows", "conflicts", "questions_for_human", "abstentions"],
    }


def build_gemini_prompt(payload: GeminiEvidenceAnalyzeRequest) -> str:
    evidence_blocks = []
    for index, item in enumerate(payload.evidence, start=1):
        text = item.text.strip()
        if len(text) > 8000:
            text = text[:8000] + "\n...[truncated]"
        source_id = item.source_id or f"source-{index:03d}"
        evidence_blocks.append(
            "\n".join(
                [
                    f"[{source_id}]",
                    f"file: {item.source_file}",
                    f"sheet_or_page: {item.source_sheet_or_page}",
                    f"location: {item.source_location}",
                    "content:",
                    text,
                ]
            )
        )
    context_lines = [f"- {key}: {value}" for key, value in payload.context.items()]
    return f"""
You are analyzing Korean weekly overtime reply materials for a reporting workflow.
Your job is not to create final report truth. Your job is to propose evidence-backed candidate rows and explicitly abstain when the evidence is weak.

Safety rules:
- Never invent a final value.
- If a number could mean headcount, hours, M/H, quantity, line count, or something else, keep it ambiguous and explain why.
- Mark anything that could cause false reporting as needs review.
- Prefer source evidence, arithmetic reconciliation, headers, row labels, sheet/file context, and known report categories.
- Output only JSON matching the schema.

Period: {payload.period_label or "(not provided)"}
Task: {payload.task}
Context:
{chr(10).join(context_lines) if context_lines else "- none"}

Evidence:
{chr(10).join(evidence_blocks)}
""".strip()


def call_gemini_json(prompt: str, schema: dict) -> dict:
    api_key = os.environ.get("GEMINI_API_KEY")
    key_status = gemini_key_status()
    if not key_status["ok"]:
        raise HTTPException(status_code=503, detail=f"GEMINI_API_KEY is not usable: {key_status['status']}")
    model = GEMINI_MODEL
    url = f"{GEMINI_API_BASE}/models/{model}:generateContent"
    request_body = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {
            "temperature": 0.1,
            "responseMimeType": "application/json",
            "responseJsonSchema": schema,
        },
    }
    data = json.dumps(request_body, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        url,
        data=data,
        method="POST",
        headers={
            "Content-Type": "application/json",
            "x-goog-api-key": api_key,
        },
    )
    try:
        with urllib.request.urlopen(request, timeout=75) as response:
            raw = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise HTTPException(status_code=502, detail=f"Gemini HTTP {exc.code}: {detail[:1000]}") from exc
    except urllib.error.URLError as exc:
        raise HTTPException(status_code=502, detail=f"Gemini request failed: {exc.reason}") from exc
    except (TimeoutError, socket.timeout) as exc:
        raise HTTPException(status_code=502, detail="Gemini request timed out. Check outbound HTTPS/firewall/proxy from the operating server.") from exc
    except OSError as exc:
        raise HTTPException(status_code=502, detail=f"Gemini network error: {type(exc).__name__}: {str(exc)[:500]}") from exc

    try:
        text = raw["candidates"][0]["content"]["parts"][0]["text"]
        parsed = json.loads(text)
    except (KeyError, IndexError, json.JSONDecodeError) as exc:
        raise HTTPException(status_code=502, detail="Gemini response did not contain valid JSON") from exc
    return {
        "ok": True,
        "model": model,
        "result": parsed,
        "usage_metadata": raw.get("usageMetadata", {}),
        "safety_note": "AI output is candidate evidence only. Final report generation requires human-confirmed locked data.",
    }


def selected_ai_evidence(run: dict, max_items: int) -> list[dict]:
    current = [item for item in run.get("evidence_items", []) if item.get("current_period")]
    pool = current or run.get("evidence_items", [])
    priority = {"xlsx_sheet": 0, "pdf_page": 1, "pptx_slide": 2}
    return sorted(
        pool,
        key=lambda item: (
            priority.get(item.get("source_type", ""), 9),
            0 if item.get("current_period") else 1,
            item.get("source_path", ""),
            item.get("source_sheet_or_page", ""),
        ),
    )[:max_items]


def run_gemini_import_batches(run: dict, period_label: str, max_ai_sources: int, batch_size: int) -> None:
    ai = run.setdefault("ai", {})
    ai.update(
        {
            "enabled": True,
            "model": GEMINI_MODEL,
            "candidate_rows": [],
            "source_assessment": [],
            "conflicts": [],
            "questions_for_human": [],
            "abstentions": [],
            "errors": [],
            "usage_metadata": [],
            "batch_count": 0,
            "evidence_sent_count": 0,
        }
    )
    evidence = selected_ai_evidence(run, max_ai_sources)
    ai["evidence_sent_count"] = len(evidence)
    context = {
        "rule_1": "AI output is candidate evidence only; never final report truth.",
        "rule_2": "관리직은 비생산으로 본다.",
        "rule_3": "직접직은 실제 생산부서만 해당한다. 생산관리/생산기술은 직접직으로 단정하지 않는다.",
        "rule_4": "source_factory와 target_factory가 불명확하면 target_factory를 비우거나 N/A로 두고 검토 필요 사유를 남긴다.",
    }
    for start in range(0, len(evidence), batch_size):
        batch = evidence[start : start + batch_size]
        ai["batch_count"] += 1
        payload = GeminiEvidenceAnalyzeRequest(
            period_label=period_label,
            context=context,
            evidence=[
                GeminiEvidenceItem(
                    source_id=item.get("source_id", ""),
                    source_file=item.get("source_file", ""),
                    source_sheet_or_page=item.get("source_sheet_or_page", ""),
                    source_location=item.get("source_location", ""),
                    text=item.get("text", ""),
                )
                for item in batch
            ],
        )
        try:
            result = call_gemini_json(build_gemini_prompt(payload), gemini_analysis_schema())
        except HTTPException as exc:
            ai["errors"].append(
                {
                    "batch": ai["batch_count"],
                    "status_code": exc.status_code,
                    "detail": exc.detail,
                    "source_ids": [item.get("source_id") for item in batch],
                }
            )
            continue
        parsed = result.get("result", {})
        for row in parsed.get("candidate_rows", []):
            ai["candidate_rows"].append({**row, "ai_batch": ai["batch_count"], "source_ids": [item.get("source_id") for item in batch]})
        ai["source_assessment"].extend(parsed.get("source_assessment", []))
        ai["conflicts"].extend(parsed.get("conflicts", []))
        ai["questions_for_human"].extend(parsed.get("questions_for_human", []))
        ai["abstentions"].extend(parsed.get("abstentions", []))
        ai["usage_metadata"].append(result.get("usage_metadata", {}))
    run["mode"] = "local_extract_plus_gemini"
    run["summary"]["ai_evidence_sent_count"] = ai["evidence_sent_count"]
    run["summary"]["ai_candidate_count"] = len(ai["candidate_rows"])
    run["summary"]["ai_error_count"] = len(ai["errors"])


def hr_connection():
    if pymysql is None:
        raise HTTPException(status_code=503, detail="pymysql dependency is not installed")
    if not hr_db_configured():
        raise HTTPException(status_code=503, detail="HR DB environment variables are not configured")
    return pymysql.connect(
        host=os.environ["HR_DB_HOST"],
        port=int(os.environ.get("HR_DB_PORT", "3306")),
        user=os.environ["HR_DB_USER"],
        password=os.environ["HR_DB_PASSWORD"],
        database=HR_DB_NAME,
        charset="utf8mb4",
        connect_timeout=5,
        read_timeout=15,
        cursorclass=pymysql.cursors.DictCursor,
    )


def production_department(text: str) -> bool:
    normalized = text.upper().replace(" ", "")
    if any(marker.upper().replace(" ", "") in normalized for marker in INDIRECT_MARKERS):
        return False
    return any(marker.upper().replace(" ", "") in normalized for marker in PRODUCTION_MARKERS)


def classify_hr_employee(row: dict) -> tuple[str, str, str]:
    dept_text = " ".join(
        str(row.get(key) or "")
        for key in ("DeptName", "emp_org_name", "org_name")
    )
    if row.get("UMJoName") == "관리직":
        return "관리직", "high", "HR UMJoName=관리직"
    if row.get("PosName") == "직접" and row.get("UMJoName") == "생산직" and production_department(dept_text):
        return "직접직", "high", "HR 직접/생산직 + 실제 생산부서"
    if row.get("PosName") == "직접" and row.get("UMJoName") == "생산직":
        return "간접직", "medium", "직접 표기는 있으나 생산부서 확증 부족"
    return "간접직", "high", "HR 간접 또는 지원부서 기준"


def entity_code_from_hr(row: dict, requested_entity_code: str | None = None) -> str:
    if requested_entity_code:
        return requested_entity_code
    binum = str(row.get("binum") or "")
    if binum == "1":
        return "daeseung"
    if binum == "2":
        return "daeseung_precision"
    if binum == "3":
        return "ilgang"
    if binum == "4":
        return "theone" if str(row.get("DeptName") or "").startswith("더원-") or row.get("PuName") == "JM서울" else "jm"
    return ""


def display_factory_from_hr(row: dict, entity_code: str) -> str:
    if entity_code == "theone":
        return "더원공장"
    if entity_code == "jm":
        return "제이엠공장"
    return str(row.get("PuName") or "")


def hr_candidate_from_row(row: dict, requested_entity_code: str | None = None) -> HREmployeeCandidate:
    entity_code = entity_code_from_hr(row, requested_entity_code)
    entity_rule = HR_ENTITY_RULES.get(entity_code, {})
    job_group, confidence, reason = classify_hr_employee(row)
    return HREmployeeCandidate(
        emp_id=str(row.get("EmpID") or ""),
        name=str(row.get("EmpName") or ""),
        entity_code=entity_code,
        entity_name=str(entity_rule.get("label") or ""),
        factory=display_factory_from_hr(row, entity_code),
        db_factory=str(row.get("PuName") or ""),
        department=str(row.get("DeptName") or ""),
        department_code=str(row.get("DeptSeq") or ""),
        org_name=str(row.get("emp_org_name") or row.get("org_name") or ""),
        org_code=str(row.get("emp_bu_code") or row.get("org_code") or ""),
        position=str(row.get("UMJpName") or ""),
        duty=str(row.get("UMJdName") or ""),
        job_family=str(row.get("UMJoName") or ""),
        direct_marker=str(row.get("PosName") or ""),
        pay_type=str(row.get("PtName") or ""),
        job_group=job_group,
        classification_reason=reason,
        confidence=confidence,
    )


def add_entity_filters(where: list[str], params: list, entity_code: str | None, factory: str | None) -> None:
    rule = HR_ENTITY_RULES.get(entity_code or "")
    if rule:
        where.append("e.binum = %s")
        params.append(rule["binum"])
        if entity_code == "theone":
            where.append("(e.PuName = %s OR e.DeptName LIKE %s)")
            params.extend(["JM서울", "더원-%"])
            return
        if entity_code == "jm":
            where.append("(e.PuName = %s OR (e.PuName LIKE %s AND e.DeptName NOT LIKE %s))")
            params.extend(["JM평택", "JM%", "더원-%"])
            return
    elif entity_code:
        where.append("1 = 0")
        return

    if factory:
        where.append("REPLACE(e.PuName, ' ', '') = REPLACE(%s, ' ', '')")
        params.append(factory)


def choose_template(kind: Literal["excel", "ppt"]) -> Path | None:
    if not TEMPLATE_DIR.exists():
        return None
    suffixes = {".xlsx"} if kind == "excel" else {".pptx"}
    files = [path for path in TEMPLATE_DIR.iterdir() if path.is_file() and path.suffix.lower() in suffixes]
    if kind == "excel":
        preferred = [path for path in files if "특근계획 검토" in path.name]
    else:
        preferred = [path for path in files if "비생산부문 주말 특근" in path.name]
    candidates = preferred or files
    return max(candidates, key=lambda path: path.stat().st_mtime) if candidates else None


def copy_template_output(kind: Literal["excel", "ppt"]) -> dict:
    template = choose_template(kind)
    if not template:
        raise HTTPException(status_code=404, detail=f"{kind} template not found in {TEMPLATE_DIR}")
    ext = template.suffix
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    generated_dir = OUTPUT_ROOT / "generated"
    generated_dir.mkdir(parents=True, exist_ok=True)
    stamped = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = generated_dir / f"{template.stem}_copy_{stamped}{ext}"
    shutil.copy2(template, output)
    return {"ok": True, "source_template": str(template), "path": str(output), "message": "원본 템플릿은 건드리지 않고 복사본을 생성했습니다."}


def build_approval_title(submission: Submission, period: Period) -> str:
    return f"{submission.factory} {period.label} 특근계획서"


def build_approval_body(submission: Submission, period: Period) -> str:
    summary = summarize(submission.rows, period)
    job_counts = split_counts(submission.rows)
    date_rows = "\n".join(
        f"<tr><td>{date}</td><td>{count}명</td></tr>"
        for date, count in sorted(summary["by_date"].items())
    )
    row_html = "\n".join(
        "<tr>"
        f"<td>{row.date}</td><td>{row.source_factory}</td><td>{row.job_group}</td>"
        f"<td>{row.team}</td><td>{row.name or '-'}</td><td>{row.position or '-'}</td>"
        f"<td>{row.headcount}명</td><td>{row.hours:g}H</td><td>{row.category1}</td><td>{row.detail}</td>"
        "</tr>"
        for row in submission.rows
    )
    attachments = "".join(f"<li>{name}</li>" for name in submission.attachments) or "<li>첨부 없음</li>"
    return f"""
<h3 style="margin:0 0 12px;">{build_approval_title(submission, period)}</h3>
<p style="margin:0 0 14px;">아래와 같이 {period.label} 특근계획을 보고드립니다.</p>
<table border="1" cellspacing="0" cellpadding="6" style="border-collapse:collapse;width:100%;margin-bottom:14px;">
  <tr><th style="width:28%;">기안부서</th><td>{submission.factory} / {submission.team}</td><th style="width:18%;">담당자</th><td>{submission.submitter}</td></tr>
  <tr><th>기간</th><td colspan="3">{period.start_date} ~ {period.end_date}</td></tr>
</table>
<table border="1" cellspacing="0" cellpadding="6" style="border-collapse:collapse;width:100%;margin-bottom:14px;">
  <tr><th>구분</th><th>인원</th><th>비고</th></tr>
  <tr><td>총 특근인원</td><td>{summary["total_headcount"]}명</td><td></td></tr>
  <tr><td>주말 특근인원</td><td>{summary["weekend_headcount"]}명</td><td>{", ".join(period.weekend_dates)}</td></tr>
  <tr><td>참고일 포함 인원</td><td>{summary["reference_headcount"]}명</td><td>{", ".join(period.reference_dates)}</td></tr>
  <tr><td>직접직/간접직/관리직</td><td colspan="2">직접 {job_counts.get("직접직", 0)}명 / 간접 {job_counts.get("간접직", 0)}명 / 관리 {job_counts.get("관리직", 0)}명</td></tr>
</table>
<table border="1" cellspacing="0" cellpadding="6" style="border-collapse:collapse;width:360px;margin-bottom:14px;">
  <tr><th>일자</th><th>인원</th></tr>
  {date_rows}
</table>
<table border="1" cellspacing="0" cellpadding="6" style="border-collapse:collapse;width:100%;margin-bottom:14px;">
  <tr><th>일자</th><th>공장</th><th>직군</th><th>팀</th><th>성명</th><th>직위</th><th>인원</th><th>시간</th><th>특근업무</th><th>세부내용</th></tr>
  {row_html}
</table>
<p style="margin:0 0 6px;"><strong>첨부</strong></p>
<ul style="margin-top:0;">{attachments}</ul>
<p style="color:#666;">※ 본문은 자동 작성 초안이며, 그룹웨어 화면에서 확인 후 수동 상신합니다.</p>
""".strip()


def clean_import_value(value) -> str:
    text = str(value or "").strip()
    return "" if text in BLANK_IMPORT_VALUES else text


def import_number(value) -> float | None:
    text = clean_import_value(value).replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return number if number >= 0 else None


def import_headcount(value) -> int | None:
    number = import_number(value)
    if number is None:
        return None
    return int(round(number))


def evidence_by_id(run: dict) -> dict[str, dict]:
    return {item.get("source_id", ""): item for item in run.get("evidence_items", [])}


def candidate_evidence_ids(row: dict) -> list[str]:
    ids: list[str] = []
    ids.extend(re.findall(r"\b(?:xlsx|xls|pdf|pptx)-\d+-\d+\b", str(row.get("evidence") or "")))
    if row.get("source_id"):
        ids.append(str(row["source_id"]))
    ids.extend(str(item) for item in row.get("source_ids") or [])
    return list(dict.fromkeys(ids))


def choose_candidate_evidence(row: dict, index: dict[str, dict]) -> dict:
    explicit_ids = re.findall(r"\b(?:xlsx|xls|pdf|pptx)-\d+-\d+\b", str(row.get("evidence") or ""))
    for source_id in explicit_ids:
        if source_id in index:
            return index[source_id]

    candidates = [index[source_id] for source_id in candidate_evidence_ids(row) if source_id in index]
    tokens = [
        token
        for token in (clean_import_value(row.get("name")), clean_import_value(row.get("detail")))
        if len(token) >= 2 and token not in {"LINE 관리", "생산 관리"}
    ]
    for item in candidates:
        text = str(item.get("text") or "")
        if any(token in text for token in tokens):
            return item

    ai_factory = clean_import_value(row.get("source_factory"))
    for item in candidates:
        evidence_factory = clean_import_value(item.get("factory_guess") or item.get("source_folder"))
        if evidence_factory == ai_factory:
            return item
    return candidates[0] if candidates else {}


def import_noise_reason(row: dict, evidence: dict | None = None) -> str:
    text = " ".join(
        clean_import_value(value)
        for value in (
            row.get("source_sheet_or_page"),
            row.get("evidence"),
            row.get("needs_review_reason"),
            evidence.get("text") if evidence else "",
        )
    )
    detail = clean_import_value(row.get("detail") or row.get("category_guess"))
    headcount = import_headcount(row.get("headcount"))
    hours = import_number(row.get("hours"))
    has_person = bool(clean_import_value(row.get("name")))
    line_plan = any(token.lower() in text.lower() for token in ("M/H", "man-hours", "라인근무계획", "수급사 및 라인근무계획"))
    if hours is not None and hours > 24 and (headcount is None or line_plan):
        return "M/H 또는 라인계획 집계값"
    if line_plan and headcount is None and not has_person and len(detail) <= 10:
        return "라인계획 보조값"
    if not detail:
        return "세부내용 없음"
    return ""


def report_category_for(row: dict) -> str:
    text = " ".join(clean_import_value(row.get(key)) for key in ("team", "detail", "category_guess", "job_group"))
    rules = [
        ("유지보수", ("유지", "보수", "보전", "설비점검", "예방보전", "교체", "수리")),
        ("생산 돌대응", ("돌발", "트러블", "긴급", "복구", "라인대응", "생산대응")),
        ("고객사 납품대응", ("고객사", "납품대응", "납품", "출하", "상차", "포장", "창고", "물류")),
        ("납품용기/제품관리", ("납품용기", "제품", "재고", "선별", "검사", "재고조사", "제품관리")),
        ("개발/공정개선", ("개발", "공정", "개선", "정도", "LAY-OUT", "OP", "TAP", "BURR")),
        ("행정업무", ("행정", "총무", "인사", "문서", "자료")),
        ("설비/시설 청소", ("청소", "정리", "이형제", "5S", "시설")),
        ("생산지원", ("생산지원", "라인 지원", "가공 생산지원")),
    ]
    for category, tokens in rules:
        if any(token.lower() in text.lower() for token in tokens):
            return category
    return "기타"


def normalize_report_category(value: object) -> str:
    text = clean_import_value(value)
    mapping = {
        "특근대응": "생산 돌대응",
        "생산특근대응": "생산 돌대응",
        "생산 돌대응": "생산 돌대응",
        "납품대응": "고객사 납품대응",
        "고객사 납품": "고객사 납품대응",
        "청소": "설비/시설 청소",
        "시설청소": "설비/시설 청소",
        "설비청소": "설비/시설 청소",
        "해외출장": "기타",
        "재물조사": "기타",
    }
    if text in REPORT_CATEGORIES:
        return text
    return mapping.get(text, text if text in REPORT_CATEGORIES else "기타")


def normalize_report_factory(entity: object, factory: object) -> str:
    entity_text = clean_import_value(entity)
    factory_text = clean_import_value(factory)
    if entity_text == "일강":
        if factory_text in {"1", "1.0", "일강1", "일강 1공장", "일강1공장"}:
            return "일강1공장"
        if factory_text in {"2", "2.0", "일강2", "일강 2공장", "일강2공장"}:
            return "일강2공장"
    if entity_text == "더원":
        return "더원공장"
    if entity_text == "제이엠":
        return "제이엠공장"
    if factory_text in REPORT_FACTORY_ORDER:
        return factory_text
    if "전산" in factory_text or "경영" in factory_text:
        return "경영지원"
    if "구매" in factory_text:
        return "구매"
    if "생기" in factory_text or "생산기술" in factory_text:
        return "생기"
    if "영업" in factory_text or "물류" in factory_text:
        return "영업"
    return factory_text or "미확인"


def normalize_report_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = clean_import_value(value)
    match = re.match(r"^(\d{2})[.\-/](\d{1,2})[.\-/](\d{1,2})", text)
    if match:
        year, month, day = match.groups()
        return f"20{year}-{int(month):02d}-{int(day):02d}"
    match = re.match(r"^(\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})", text)
    if match:
        year, month, day = match.groups()
        return f"{year}-{int(month):02d}-{int(day):02d}"
    return text


def reference_workbook_path() -> Path | None:
    if REFERENCE_WORKBOOK_ENV:
        path = Path(REFERENCE_WORKBOOK_ENV)
        return path if path.exists() else None
    if TEMPLATE_DIR.exists():
        matches = sorted(
            path for path in TEMPLATE_DIR.glob("*특근계획 검토*.xlsx")
            if not path.name.startswith("~$")
        )
        if matches:
            return matches[0]
    return None


def column_index(header: dict[str, int], *names: str, fallback: int) -> int:
    for name in names:
        if name in header:
            return header[name]
    return fallback


def reference_report_baseline(period: Period) -> dict:
    static_baseline = REPORT_BASELINES.get(period.id, {})
    fallback = {
        **static_baseline,
        "source": "static_baseline",
        "row_count": static_baseline.get("row_count", 0),
        "by_category": {
            category: static_baseline.get("by_category", {}).get(category, 0)
            for category in REPORT_CATEGORIES
        },
        "by_factory": dict(static_baseline.get("by_factory", {})),
        "matrix": {},
    }
    path = reference_workbook_path()
    if openpyxl is None or not path:
        return fallback

    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet_name = "통합(5월 5주차)" if "통합(5월 5주차)" in workbook.sheetnames else workbook.sheetnames[0]
        sheet = workbook[sheet_name]
        header_row = 4
        for row_number in range(1, min(sheet.max_row, 30) + 1):
            values = [clean_import_value(sheet.cell(row_number, col).value) for col in range(1, min(sheet.max_column, 50) + 1)]
            if "일자" in values and "인원" in values and "특근내용(분류1)" in values:
                header_row = row_number
                break
        header = {
            clean_import_value(sheet.cell(header_row, col).value): col
            for col in range(1, sheet.max_column + 1)
            if clean_import_value(sheet.cell(header_row, col).value)
        }
        no_col = column_index(header, "NO", fallback=2)
        date_col = column_index(header, "일자", fallback=3)
        entity_col = column_index(header, "소속", fallback=7)
        factory_col = column_index(header, "공장", fallback=8)
        headcount_col = column_index(header, "인원", fallback=13)
        category_col = column_index(header, "특근내용(분류1)", fallback=14)

        rows: list[dict] = []
        by_date: defaultdict[str, int] = defaultdict(int)
        by_category: defaultdict[str, int] = defaultdict(int)
        by_factory: defaultdict[str, int] = defaultdict(int)
        matrix: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
        for row_number in range(header_row + 1, sheet.max_row + 1):
            no_value = sheet.cell(row_number, no_col).value
            count = import_headcount(sheet.cell(row_number, headcount_col).value)
            date = normalize_report_date(sheet.cell(row_number, date_col).value)
            if count is None or not date:
                continue
            if not isinstance(no_value, (int, float)) and not clean_import_value(no_value).isdigit():
                continue
            category = normalize_report_category(sheet.cell(row_number, category_col).value)
            factory = normalize_report_factory(sheet.cell(row_number, entity_col).value, sheet.cell(row_number, factory_col).value)
            rows.append({"date": date, "category": category, "factory": factory, "headcount": count})
            by_date[date] += count
            by_category[category] += count
            by_factory[factory] += count
            matrix[category][factory] += count

        total = sum(by_date.values())
        weekend = sum(value for date, value in by_date.items() if date in period.weekend_dates)
        reference = sum(value for date, value in by_date.items() if date in period.reference_dates)
        return {
            "source": "reference_workbook",
            "workbook_path": str(path),
            "sheet_name": sheet_name,
            "row_count": len(rows),
            "total_headcount": total,
            "weekend_headcount": weekend,
            "reference_headcount": reference,
            "by_date": dict(sorted(by_date.items())),
            "by_category": {category: by_category.get(category, 0) for category in REPORT_CATEGORIES},
            "by_factory": dict(sorted(by_factory.items(), key=lambda item: (-item[1], item[0]))),
            "matrix": {category: dict(factory_map) for category, factory_map in matrix.items()},
        }
    except Exception as exc:  # pragma: no cover - defensive fallback for damaged workbooks.
        return {**fallback, "warning": f"reference workbook read failed: {type(exc).__name__}: {exc}"}


def import_report_rows(run: dict, period: Period) -> tuple[list[dict], list[dict], list[dict]]:
    index = evidence_by_id(run)
    rows: list[dict] = []
    exclusions: list[dict] = []
    exceptions: list[dict] = []
    seen: set[tuple] = set()

    for raw in run.get("ai", {}).get("candidate_rows", []) + run.get("local_candidate_rows", []):
        evidence = choose_candidate_evidence(raw, index)
        source_id = evidence.get("source_id") or (candidate_evidence_ids(raw) or [""])[0]
        factory = clean_import_value(evidence.get("factory_guess") or evidence.get("source_folder") or raw.get("source_factory"))
        date = clean_import_value(raw.get("date"))
        name = clean_import_value(raw.get("name"))
        team = clean_import_value(raw.get("team"))
        detail = clean_import_value(raw.get("detail") or raw.get("category_guess"))
        headcount = import_headcount(raw.get("headcount"))
        hours = import_number(raw.get("hours"))
        if headcount is None and name:
            headcount = 1
        noise = import_noise_reason(raw, evidence)
        key = (date, factory, team, name, headcount, hours, detail, source_id)
        if key in seen:
            continue
        seen.add(key)
        base_row = {
            "source_id": source_id,
            "source_file": evidence.get("source_file") or raw.get("source_file") or "",
            "source_sheet_or_page": evidence.get("source_sheet_or_page") or raw.get("source_sheet_or_page") or "",
            "factory": factory,
            "date": date,
            "team": team,
            "name": name,
            "headcount": headcount,
            "hours": hours,
            "detail": detail,
            "confidence": clean_import_value(raw.get("confidence")) or "Low",
            "category": report_category_for(raw),
        }
        missing = []
        if date not in period.weekend_dates + period.reference_dates:
            missing.append("일자")
        if not factory:
            missing.append("공장")
        if not detail:
            missing.append("세부내용")
        if not headcount:
            missing.append("인원")
        if noise:
            exclusions.append({**base_row, "reason": noise})
            continue
        if missing:
            exceptions.append({**base_row, "severity": "high", "reason": f"필수값 누락: {', '.join(missing)}"})
            continue
        rows.append(base_row)

    return rows, exclusions, exceptions


def summarize_report_rows(rows: list[dict], period: Period) -> dict:
    by_date: defaultdict[str, int] = defaultdict(int)
    by_category: defaultdict[str, int] = defaultdict(int)
    by_factory: defaultdict[str, int] = defaultdict(int)
    matrix: dict[str, dict[str, dict[str, int]]] = {
        category: {factory: {"total": 0, "weekend": 0} for factory in REPORT_FACTORY_ORDER}
        for category in REPORT_CATEGORIES
    }
    extra_factories: list[str] = []

    for row in rows:
        count = int(row.get("headcount") or 0)
        date = row.get("date") or ""
        category = row.get("category") if row.get("category") in REPORT_CATEGORIES else "기타"
        factory = row.get("factory") or "미확인"
        if factory not in REPORT_FACTORY_ORDER and factory not in extra_factories:
            extra_factories.append(factory)
            for category_map in matrix.values():
                category_map[factory] = {"total": 0, "weekend": 0}
        by_date[date] += count
        by_category[category] += count
        by_factory[factory] += count
        matrix[category][factory]["total"] += count
        if date in period.weekend_dates:
            matrix[category][factory]["weekend"] += count

    total = sum(by_date.values())
    weekend = sum(value for date, value in by_date.items() if date in period.weekend_dates)
    reference = sum(value for date, value in by_date.items() if date in period.reference_dates)
    return {
        "total_headcount": total,
        "weekend_headcount": weekend,
        "reference_headcount": reference,
        "by_date": dict(sorted(by_date.items())),
        "by_category": {category: by_category.get(category, 0) for category in REPORT_CATEGORIES},
        "by_factory": dict(sorted(by_factory.items(), key=lambda item: (-item[1], item[0]))),
        "factory_order": [*REPORT_FACTORY_ORDER, *extra_factories],
        "matrix": matrix,
    }


def reconciliation_line(label: str, expected: int | None, actual: int | None, unit: str = "명") -> dict:
    expected_value = int(expected or 0)
    actual_value = int(actual or 0)
    delta = actual_value - expected_value
    if delta == 0:
        status = "match"
    elif delta > 0:
        status = "over"
    else:
        status = "missing"
    return {
        "label": label,
        "expected": expected_value,
        "actual": actual_value,
        "delta": delta,
        "unit": unit,
        "status": status,
    }


def source_label(item: dict) -> str:
    file_name = Path(clean_import_value(item.get("source_file"))).name if clean_import_value(item.get("source_file")) else ""
    sheet = clean_import_value(item.get("source_sheet_or_page"))
    if file_name and sheet:
        return f"{file_name} / {sheet}"
    return file_name or sheet or clean_import_value(item.get("source_id")) or "미확인 원천"


def build_source_coverage(run: dict, rows: list[dict], exclusions: list[dict], row_exceptions: list[dict]) -> list[dict]:
    coverage: dict[str, dict] = {}

    def ensure(label: str, item: dict | None = None) -> dict:
        if label not in coverage:
            coverage[label] = {
                "label": label,
                "source_file": clean_import_value((item or {}).get("source_file")),
                "source_sheet_or_page": clean_import_value((item or {}).get("source_sheet_or_page")),
                "counted_headcount": 0,
                "counted_rows": 0,
                "excluded_rows": 0,
                "exception_rows": 0,
                "ai_error_rows": 0,
                "notes": [],
            }
        return coverage[label]

    for row in rows:
        entry = ensure(source_label(row), row)
        entry["counted_headcount"] += int(row.get("headcount") or 0)
        entry["counted_rows"] += 1
    for row in exclusions:
        entry = ensure(source_label(row), row)
        entry["excluded_rows"] += 1
        reason = clean_import_value(row.get("reason"))
        if reason and reason not in entry["notes"]:
            entry["notes"].append(reason)
    for row in row_exceptions:
        entry = ensure(source_label(row), row)
        entry["exception_rows"] += 1
        reason = clean_import_value(row.get("reason"))
        if reason and reason not in entry["notes"]:
            entry["notes"].append(reason)

    index = evidence_by_id(run)
    for error in run.get("ai", {}).get("errors", []):
        for source_id in error.get("source_ids") or []:
            evidence = index.get(str(source_id), {"source_id": str(source_id)})
            entry = ensure(source_label(evidence), evidence)
            entry["ai_error_rows"] += 1
            detail = clean_import_value(error.get("detail"))
            if detail and detail not in entry["notes"]:
                entry["notes"].append(detail)

    return sorted(
        coverage.values(),
        key=lambda item: (
            -(item["exception_rows"] + item["ai_error_rows"] + item["excluded_rows"]),
            -item["counted_headcount"],
            item["label"],
        ),
    )[:40]


def build_reconciliation(run: dict, period: Period, rows: list[dict], exclusions: list[dict], row_exceptions: list[dict], summary: dict, baseline: dict) -> dict:
    totals = [
        reconciliation_line("전체 특근인원", baseline.get("total_headcount"), summary.get("total_headcount")),
        reconciliation_line("주말 특근현황", baseline.get("weekend_headcount"), summary.get("weekend_headcount")),
        reconciliation_line("6/3 포함 참고", baseline.get("reference_headcount"), summary.get("reference_headcount")),
    ]

    date_keys = list(dict.fromkeys([*period.weekend_dates, *period.reference_dates, *baseline.get("by_date", {}).keys(), *summary.get("by_date", {}).keys()]))
    dates = [
        reconciliation_line(date, baseline.get("by_date", {}).get(date), summary.get("by_date", {}).get(date))
        for date in date_keys
    ]

    category_keys = list(dict.fromkeys([*REPORT_CATEGORIES, *baseline.get("by_category", {}).keys(), *summary.get("by_category", {}).keys()]))
    categories = [
        reconciliation_line(category, baseline.get("by_category", {}).get(category), summary.get("by_category", {}).get(category))
        for category in category_keys
        if baseline.get("by_category", {}).get(category) or summary.get("by_category", {}).get(category)
    ]

    factory_keys = list(dict.fromkeys([*REPORT_FACTORY_ORDER, *baseline.get("by_factory", {}).keys(), *summary.get("by_factory", {}).keys()]))
    factories = [
        reconciliation_line(factory, baseline.get("by_factory", {}).get(factory), summary.get("by_factory", {}).get(factory))
        for factory in factory_keys
        if baseline.get("by_factory", {}).get(factory) or summary.get("by_factory", {}).get(factory)
    ]

    drivers: list[dict] = []
    for line in [*totals, *dates]:
        if line["delta"]:
            direction = "과다" if line["delta"] > 0 else "부족"
            drivers.append({
                "severity": "high",
                "label": line["label"],
                "message": f"{abs(line['delta'])}{line['unit']} {direction}",
            })
    for line in sorted([*categories, *factories], key=lambda item: abs(item["delta"]), reverse=True)[:10]:
        if line["delta"]:
            direction = "과다" if line["delta"] > 0 else "부족"
            drivers.append({
                "severity": "medium",
                "label": line["label"],
                "message": f"{abs(line['delta'])}{line['unit']} {direction}",
            })
    if exclusions:
        drivers.append({"severity": "medium", "label": "집계 제외", "message": f"{len(exclusions)}건은 M/H/라인계획/빈 세부내용으로 제외"})
    if run.get("ai", {}).get("errors"):
        drivers.append({"severity": "high", "label": "AI 미처리", "message": f"{len(run.get('ai', {}).get('errors', []))}개 배치/원천 재처리 필요"})

    return {
        "baseline_source": baseline.get("source"),
        "baseline_workbook": baseline.get("workbook_path", ""),
        "baseline_sheet": baseline.get("sheet_name", ""),
        "baseline_warning": baseline.get("warning", ""),
        "totals": totals,
        "dates": dates,
        "categories": categories,
        "factories": factories,
        "drivers": drivers[:18],
        "source_coverage": build_source_coverage(run, rows, exclusions, row_exceptions),
    }


def report_exceptions(run: dict, period: Period, rows: list[dict], exclusions: list[dict], row_exceptions: list[dict], summary: dict, baseline: dict) -> list[dict]:
    exceptions = list(row_exceptions)
    for key, label in (
        ("total_headcount", "전체 특근인원"),
        ("weekend_headcount", "주말 특근현황"),
        ("reference_headcount", "6/3 포함 참고"),
    ):
        expected = baseline.get(key)
        actual = summary.get(key)
        if expected is not None and expected != actual:
            exceptions.insert(0, {"severity": "high", "reason": f"{label} 기준 불일치: 기준 {expected}명 / 현재 초안 {actual}명"})

    if exclusions:
        exceptions.append({"severity": "medium", "reason": f"집계 제외 후보 {len(exclusions)}건: M/H, 라인계획, 세부내용 없음 등"})
    for error in run.get("ai", {}).get("errors", []):
        source_ids = ", ".join(str(item) for item in error.get("source_ids") or [])
        exceptions.append({"severity": "high", "reason": f"AI 미처리 원천 {source_ids}: {error.get('detail', '')}"})

    people: defaultdict[str, set[str]] = defaultdict(set)
    for row in rows:
        name = clean_import_value(row.get("name"))
        if name and "외" not in name:
            people[name].add(row.get("factory") or "")
    for name, factories in sorted(people.items()):
        if len(factories) > 1:
            exceptions.append({"severity": "medium", "reason": f"동명이인/공장충돌 가능: {name} -> {', '.join(sorted(factories))}"})
    return exceptions[:80]


def build_ppt_report_workspace(period: Period) -> dict:
    baseline = reference_report_baseline(period)
    try:
        run = load_latest(IMPORT_RUN_DIR)
    except (FileNotFoundError, json.JSONDecodeError) as exc:
        empty_summary = summarize_report_rows([], period)
        reconciliation = build_reconciliation({}, period, [], [], [], empty_summary, baseline)
        return {
            "source": "seed_state",
            "ready": False,
            "message": f"import run not readable: {IMPORT_RUN_DIR} ({type(exc).__name__})",
            "rows": [],
            "exclusions": [],
            "exceptions": [{"severity": "high", "reason": "회신자료 import 결과가 없어 PPT 자동 산출을 시작할 수 없음"}],
            "summary": empty_summary,
            "baseline": baseline,
            "reconciliation": reconciliation,
            "counts": {"input_candidates": 0, "report_rows": 0, "excluded_rows": 0, "exception_count": 1, "blocking_count": 1},
        }
    if not isinstance(run, dict):
        empty_summary = summarize_report_rows([], period)
        reconciliation = build_reconciliation({}, period, [], [], [], empty_summary, baseline)
        return {
            "source": "seed_state",
            "ready": False,
            "message": f"import run result not found: {IMPORT_RUN_DIR}",
            "rows": [],
            "exclusions": [],
            "exceptions": [{"severity": "high", "reason": "회신자료 import 결과가 없거나 latest.json이 비어 있음"}],
            "summary": empty_summary,
            "baseline": baseline,
            "reconciliation": reconciliation,
            "counts": {"input_candidates": 0, "report_rows": 0, "excluded_rows": 0, "exception_count": 1, "blocking_count": 1},
        }

    rows, exclusions, row_exceptions = import_report_rows(run, period)
    summary = summarize_report_rows(rows, period)
    reconciliation = build_reconciliation(run, period, rows, exclusions, row_exceptions, summary, baseline)
    exceptions = report_exceptions(run, period, rows, exclusions, row_exceptions, summary, baseline)
    blocking = [item for item in exceptions if item.get("severity") == "high"]
    return {
        "source": "latest_import",
        "ready": not blocking,
        "message": "PPT 1~3페이지 산출 가능" if not blocking else "PPT 잠금 전 예외 확인 필요",
        "run_id": run.get("run_id"),
        "rows": rows,
        "exclusions": exclusions[:30],
        "exceptions": exceptions,
        "summary": summary,
        "baseline": baseline,
        "reconciliation": reconciliation,
        "counts": {
            "input_candidates": len(run.get("ai", {}).get("candidate_rows", [])) + len(run.get("local_candidate_rows", [])),
            "report_rows": len(rows),
            "excluded_rows": len(exclusions),
            "exception_count": len(exceptions),
            "blocking_count": len(blocking),
        },
    }


app = FastAPI(title="특근 보고 취합 WEB", version=APP_VERSION)
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


@app.get("/")
def index() -> FileResponse:
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/health")
def health() -> dict:
    return {
        "ok": True,
        "version": APP_VERSION,
        "service": "overtime-reporting-web",
        "features": APP_FEATURES,
    }


@app.get("/api/version")
def version() -> dict:
    return {"product": "특근 보고 취합 WEB", "version": APP_VERSION, "features": APP_FEATURES}


@app.get("/api/bootstrap")
def bootstrap() -> dict:
    state = load_state()
    period = state.periods[0] if state.periods else None
    return {
        "periods": [p.model_dump() for p in state.periods],
        "legal_entities": [e.model_dump() for e in state.legal_entities],
        "submissions": [s.model_dump() for s in state.submissions],
        "dashboard": summarize(all_rows(state), period),
        "config": config(),
    }


@app.get("/api/config")
def config() -> dict:
    return {
        "template_dir": str(TEMPLATE_DIR),
        "template_dir_exists": TEMPLATE_DIR.exists(),
        "reference_workbook": str(reference_workbook_path() or ""),
        "reference_workbook_exists": bool(reference_workbook_path()),
        "output_root": str(OUTPUT_ROOT),
        "reply_import": {
            "default_reply_dir": str(REPLY_DIR),
            "default_reply_dir_exists": REPLY_DIR.exists(),
            "import_run_dir": str(IMPORT_RUN_DIR),
            "policy": "local_extract_first_gemini_requires_explicit_allow",
        },
        "hr_db": {
            "host": os.environ.get("HR_DB_HOST", ""),
            "port": os.environ.get("HR_DB_PORT", "3306"),
            "database": HR_DB_NAME,
            "user_configured": bool(os.environ.get("HR_DB_USER")),
            "password_configured": bool(os.environ.get("HR_DB_PASSWORD")),
            "database_configured": bool(HR_DB_NAME),
            "lookup_available": bool(hr_db_configured() and pymysql is not None),
            "policy": "read_only_select_only",
        },
        "hr_factory_options": {
            code: values["factories"]
            for code, values in HR_ENTITY_RULES.items()
        },
        "ai": {
            "provider": "gemini",
            "model": GEMINI_MODEL,
            "configured": gemini_configured(),
            "key_status": gemini_key_status()["status"],
            "policy": "candidate_evidence_only_no_final_report_without_human_lock",
        },
    }


@app.get("/api/ai/gemini/status")
def gemini_status() -> dict:
    key_status = gemini_key_status()
    return {
        "configured": gemini_configured(),
        "provider": "gemini",
        "model": GEMINI_MODEL,
        "api_base": GEMINI_API_BASE,
        "key_status": key_status["status"],
        "policy": "candidate_evidence_only_no_final_report_without_human_lock",
    }


@app.post("/api/ai/gemini/analyze-evidence")
def analyze_evidence_with_gemini(payload: GeminiEvidenceAnalyzeRequest) -> dict:
    if not payload.evidence:
        raise HTTPException(status_code=400, detail="evidence is required")
    try:
        prompt = build_gemini_prompt(payload)
        return call_gemini_json(prompt, gemini_analysis_schema())
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Gemini analysis failed: {type(exc).__name__}: {str(exc)[:500]}") from exc


@app.get("/api/import/latest")
def import_latest() -> dict:
    latest = load_latest(IMPORT_RUN_DIR)
    if not latest:
        return {
            "ok": False,
            "message": "import run result not found",
            "import_run_dir": str(IMPORT_RUN_DIR),
            "default_reply_dir": str(REPLY_DIR),
        }
    return {"ok": True, "run": latest, "import_run_dir": str(IMPORT_RUN_DIR)}


@app.post("/api/import/run")
def import_run(payload: ReplyImportRunRequest) -> dict:
    reply_dir = Path(payload.reply_dir).expanduser() if payload.reply_dir else REPLY_DIR
    if not reply_dir.exists():
        raise HTTPException(status_code=404, detail=f"reply_dir not found: {reply_dir}")
    run = run_reply_import(reply_dir, payload.period_label)
    if payload.allow_gemini:
        if not gemini_configured():
            raise HTTPException(status_code=503, detail="GEMINI_API_KEY environment variable is not configured")
        run_gemini_import_batches(run, payload.period_label, payload.max_ai_sources, payload.ai_batch_size)
    paths = save_run(run, IMPORT_RUN_DIR)
    return {"ok": True, "run": run, "paths": paths}


@app.get("/api/hr/status")
def hr_status() -> dict:
    return {
        "configured": hr_db_configured(),
        "dependency_available": pymysql is not None,
        "database": HR_DB_NAME,
        "policy": "read_only_select_only",
        "factory_options": {code: values["factories"] for code, values in HR_ENTITY_RULES.items()},
    }


@app.get("/api/hr/employees/search")
def search_hr_employees(
    q: str = Query("", min_length=0, max_length=40),
    entity_code: str | None = Query(None, max_length=40),
    factory: str | None = Query(None, max_length=40),
    limit: int = Query(20, ge=1, le=50),
) -> dict:
    query = q.strip()
    if len(query) < 2:
        return {"ok": True, "items": [], "count": 0, "message": "이름을 2글자 이상 입력하세요."}

    where = ["e.TypeName = %s", "e.binum IN ('1', '2', '3', '4')", "e.EmpName LIKE %s"]
    params: list = ["재직자", f"%{query}%"]
    add_entity_filters(where, params, entity_code, factory)
    params.extend([query, limit])
    sql = f"""
        SELECT
            e.EmpName,
            e.EmpID,
            e.binum,
            e.PuName,
            e.DeptName,
            e.DeptSeq,
            e.PosName,
            e.UMJpName,
            e.UMJdName,
            e.UMJoName,
            e.PtName,
            e.TypeName,
            e.TypeSeq,
            e.n_bu_name AS emp_org_name,
            e.bu_code AS emp_bu_code,
            e.up_bu_code AS emp_up_bu_code,
            b.n_bu_name AS org_name,
            b.bu_code AS org_code,
            b.up_bu_code AS org_parent_code,
            b.lv_no AS org_level
        FROM ds_t_emp e
        LEFT JOIN (
            SELECT
                old_bi_code,
                old_bu_code,
                MIN(n_bu_name) AS n_bu_name,
                MIN(bu_code) AS bu_code,
                MIN(up_bu_code) AS up_bu_code,
                MIN(lv_no) AS lv_no
            FROM buseo_t
            WHERE old_bi_code IS NOT NULL AND old_bu_code IS NOT NULL
            GROUP BY old_bi_code, old_bu_code
        ) b
            ON e.DeptSeq = CAST(b.old_bu_code AS CHAR)
            AND e.binum = CAST(b.old_bi_code AS CHAR)
        WHERE {" AND ".join(where)}
        ORDER BY
            CASE WHEN e.EmpName = %s THEN 0 ELSE 1 END,
            e.PuName,
            e.DeptName,
            e.EmpName,
            e.EmpID
        LIMIT %s
    """
    with hr_connection() as conn:
        with conn.cursor() as cur:
            cur.execute("SET SESSION TRANSACTION READ ONLY")
            cur.execute("START TRANSACTION READ ONLY")
            cur.execute(sql, params)
            rows = cur.fetchall()
            cur.execute("ROLLBACK")
    items = [hr_candidate_from_row(row, entity_code).model_dump() for row in rows]
    return {
        "ok": True,
        "items": items,
        "count": len(items),
        "policy": "read_only_select_only",
    }


@app.get("/api/templates")
def templates() -> dict:
    files = []
    if TEMPLATE_DIR.exists():
        files = [
            {"name": path.name, "path": str(path), "size": path.stat().st_size}
            for path in sorted(TEMPLATE_DIR.iterdir())
            if path.is_file() and path.suffix.lower() in {".xlsx", ".xls", ".pptx"}
        ]
    return {"template_dir": str(TEMPLATE_DIR), "files": files}


@app.get("/api/dashboard")
def dashboard() -> dict:
    state = load_state()
    period = state.periods[0] if state.periods else None
    return summarize(all_rows(state), period)


@app.get("/api/submissions")
def submissions() -> list[dict]:
    return [s.model_dump() for s in load_state().submissions]


@app.post("/api/submissions")
def create_submission(payload: SubmissionCreate) -> dict:
    state = load_state()
    period = state.periods[0]
    new_id = f"sub-{len(state.submissions) + 1:03d}"
    submission = Submission(
        id=new_id,
        period_id=period.id,
        submitter=payload.submitter,
        entity_code=payload.entity_code,
        factory=payload.factory,
        team=payload.team,
        rows=payload.rows,
        updated_at=_now(),
    )
    state.submissions.append(submission)
    save_state(state)
    return submission.model_dump()


@app.get("/api/report/preview")
def report_preview() -> dict:
    state = load_state()
    period = state.periods[0]
    rows = all_rows(state)
    summary = summarize(rows, period)
    date_items = [{"label": date, "value": value} for date, value in sorted(summary["by_date"].items())]
    category_items = [{"label": key, "value": value} for key, value in sorted(summary["by_category"].items(), key=lambda item: item[1], reverse=True)]
    factory_items = [{"label": key, "value": value} for key, value in sorted(summary["by_factory"].items(), key=lambda item: item[1], reverse=True)]
    return {
        "period": period.model_dump(),
        "summary": summary,
        "primary_goal": "보고자료 PPT 1~3페이지 완성",
        "slides": [
            {
                "slide": 1,
                "title": "표지",
                "items": [
                    {"label": "보고명", "value": f"비생산부문 주말 특근현황 보고 [ {period.label.replace('2026년 ', '')} ]"},
                    {"label": "기준기간", "value": f"{period.start_date} ~ {period.end_date}"},
                    {"label": "작성부서", "value": "경영기획본부"},
                ],
            },
            {
                "slide": 2,
                "title": "비생산부문 주말 특근 종합현황",
                "items": [
                    {"label": "전체 특근인원", "value": summary["total_headcount"]},
                    {"label": "주말 특근현황", "value": summary["weekend_headcount"]},
                    {"label": "6/3 포함 참고", "value": summary["reference_headcount"]},
                    *date_items,
                    *category_items[:10],
                ],
            },
            {
                "slide": 3,
                "title": "공장별 비생산부문 특근업무 세부",
                "items": [
                    *factory_items,
                    *category_items[:10],
                ],
            },
        ],
    }


@app.get("/api/report/workspace")
def report_workspace() -> dict:
    state = load_state()
    period = state.periods[0]
    report = build_ppt_report_workspace(period)
    summary = report["summary"]
    date_items = [{"label": date, "value": value} for date, value in sorted(summary["by_date"].items())]
    category_items = [{"label": key, "value": value} for key, value in summary["by_category"].items() if value]
    factory_items = [{"label": key, "value": value} for key, value in summary["by_factory"].items()]
    return {
        "period": period.model_dump(),
        "primary_goal": "보고자료 PPT 1~3페이지 완성",
        "report_draft": report,
        "slides": [
            {
                "slide": 1,
                "title": "표지",
                "items": [
                    {"label": "보고명", "value": f"비생산부문 주말 특근현황 보고 [ {period.label} ]"},
                    {"label": "기준기간", "value": f"{period.start_date} ~ {period.end_date}"},
                    {"label": "작성부서", "value": "경영기획본부"},
                ],
            },
            {
                "slide": 2,
                "title": "비생산부문 주말 특근 종합현황",
                "items": [
                    {"label": "전체 특근인원", "value": summary["total_headcount"]},
                    {"label": "주말 특근현황", "value": summary["weekend_headcount"]},
                    {"label": "6/3 포함 참고", "value": summary["reference_headcount"]},
                    *date_items,
                    *category_items[:10],
                ],
            },
            {
                "slide": 3,
                "title": "공장별 비생산부문 특근업무 세부",
                "items": [
                    *factory_items,
                    *category_items[:10],
                ],
            },
        ],
    }


@app.post("/api/report/export-excel")
def export_excel_stub() -> dict:
    return copy_template_output("excel")


@app.post("/api/report/export-ppt")
def export_ppt_stub() -> dict:
    return copy_template_output("ppt")


@app.get("/api/overtime-submissions/{submission_id}/approval-preview")
def approval_preview(submission_id: str) -> dict:
    state = load_state()
    period = state.periods[0]
    submission = next((s for s in state.submissions if s.id == submission_id), None)
    if not submission:
        raise HTTPException(status_code=404, detail="submission not found")
    entity = next((e for e in state.legal_entities if e.entity_code == submission.entity_code), None)
    return {
        "submission_id": submission_id,
        "title": build_approval_title(submission, period),
        "body_html": build_approval_body(submission, period),
        "form_label": entity.default_overtime_form_label if entity else "",
        "submit_policy": "draft_only_never_auto_submit",
    }


@app.post("/api/overtime-submissions/{submission_id}/fill-approval-draft")
def fill_approval_draft_stub(submission_id: str) -> dict:
    state = load_state()
    submission = next((s for s in state.submissions if s.id == submission_id), None)
    if not submission:
        raise HTTPException(status_code=404, detail="submission not found")
    submission.groupware_draft_status = "adapter_pending"
    submission.updated_at = _now()
    save_state(state)
    return {
        "ok": False,
        "raw_status": "adapter_pending",
        "message": "그룹웨어 브라우저 자동입력 어댑터는 아직 연결 전입니다. 이 API는 상신을 수행하지 않습니다.",
        "submission": submission.model_dump(),
    }
