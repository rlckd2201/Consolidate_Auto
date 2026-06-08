from __future__ import annotations

import json
import re
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET


TARGET_DATE_TOKENS = (
    "2026-05-30",
    "2026-05-31",
    "2026-06-03",
    "05.30",
    "05.31",
    "06.03",
    "5.30",
    "5.31",
    "6.3",
    "5/30",
    "5/31",
    "6/3",
    "5월 30일",
    "5월 31일",
    "6월 3일",
)

TARGET_ISO_DATES = {"2026-05-30", "2026-05-31", "2026-06-03"}

KEYWORDS = (
    "특근",
    "특근인원",
    "인원",
    "시간",
    "M/H",
    "상세",
    "세부",
    "합계",
    "TOTAL",
    "관리직",
    "간접",
    "직접",
    "비생산",
    "생산",
    "보전",
    "물류",
    "품질",
    "설비",
    "납품",
    "유지보수",
    "재물",
    "출장",
    "청소",
    "지원",
)

CATEGORY_RULES = (
    ("유지보수", ("보전", "설비점검", "예방보전", "유지보수", "수리")),
    ("납품대응", ("납품", "물류", "출하", "창고", "고객사")),
    ("납품용기/제품관리", ("용기", "제품", "검사", "선별", "반송품", "제품조사")),
    ("설비/시설 청소", ("청소", "이형제", "시설")),
    ("행정업무", ("총무", "인사", "노사", "행정")),
    ("생산지원", ("생산지원", "라인지원", "가공라인", "주조라인")),
    ("기타", ("재물", "출장", "실사")),
)

JOB_GROUPS = ("관리직", "간접직", "직접직", "관리", "간접", "직접")
LINE_SUMMARY_NAMES = {"", "주간", "야간", "계", "합계", "직접직", "간접직", "관리직"}
STRUCTURED_REPORTABLE_TEXT = (
    "관리",
    "총괄",
    "전산",
    "ISIR",
    "지적",
    "집계",
    "월마감",
    "서류",
    "자료",
    "재물",
    "지원판단",
    "출하대응",
    "공정관리",
    "품질문제",
    "개선",
    "개조",
    "교체",
    "수정",
    "점검",
    "설비",
    "유지보수",
    "청소",
    "선별",
    "불량",
    "반송품",
    "포장",
    "납품",
    "출하",
    "상차",
    "불출",
    "유류",
)


class SnapshotCell:
    def __init__(self, value: Any):
        self.value = value


class WorksheetSnapshot:
    def __init__(self, title: str, rows: list[list[Any]], max_column: int):
        self.title = title
        self._rows = rows
        self.max_row = len(rows)
        self.max_column = max_column

    def cell(self, row: int, column: int) -> SnapshotCell:
        value = None
        if 1 <= row <= len(self._rows):
            values = self._rows[row - 1]
            if 1 <= column <= len(values):
                value = values[column - 1]
        return SnapshotCell(value)

def now_id() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def norm(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if hasattr(value, "strftime"):
        try:
            return value.strftime("%Y-%m-%d")
        except Exception:
            pass
    return str(value).replace("\n", " ").strip()


def rel(path: Path, root: Path) -> str:
    try:
        return str(path.relative_to(root))
    except ValueError:
        return str(path)


def row_to_text(values: list[Any]) -> str:
    return " / ".join(text for text in (norm(value) for value in values) if text)


def contains_any(text: str, tokens: tuple[str, ...]) -> bool:
    lowered = text.lower()
    return any(token.lower() in lowered for token in tokens)


def hits(text: str, tokens: tuple[str, ...]) -> list[str]:
    lowered = text.lower()
    return [token for token in tokens if token.lower() in lowered]


def role_guess(path: Path) -> str:
    name = path.name
    ext = path.suffix.lower()
    if ext == ".xls":
        return "legacy_xls_needs_conversion"
    if "52" in name or "근무시간" in name:
        return "work_hour_support"
    if "신고서" in name or "특근계획" in name:
        return "plan_declaration"
    if "상세" in name or "현황" in name or ext == ".xlsx":
        return "structured_candidate_source"
    if ext == ".pdf":
        return "approval_pdf_evidence"
    if ext == ".pptx":
        return "presentation_reference"
    return "unknown"


def extract_factory(text: str, fallback: str = "") -> str:
    for factory in ("D1공장", "D2공장", "D3공장", "P1공장", "P2공장", "P3공장", "P4공장", "일강1공장", "일강 1공장", "일강2공장", "일강 2공장", "더원공장", "제이엠공장"):
        if factory in text:
            return factory.replace("일강 1공장", "일강1공장").replace("일강 2공장", "일강2공장")
    for factory in ("1공장", "2공장"):
        if "일강" in text and factory in text:
            return f"일강{factory}"
    return fallback


def extract_date(text: str) -> str:
    explicit = re.search(r"20\d{2}[-./]\d{1,2}[-./]\d{1,2}", text)
    if explicit:
        parts = re.split(r"[-./]", explicit.group(0))
        return f"{int(parts[0]):04d}-{int(parts[1]):02d}-{int(parts[2]):02d}"
    mapped = {
        "5/30": "2026-05-30",
        "5.30": "2026-05-30",
        "05.30": "2026-05-30",
        "5월 30일": "2026-05-30",
        "5/31": "2026-05-31",
        "5.31": "2026-05-31",
        "05.31": "2026-05-31",
        "5월 31일": "2026-05-31",
        "6/3": "2026-06-03",
        "6.3": "2026-06-03",
        "06.03": "2026-06-03",
        "6월 3일": "2026-06-03",
    }
    for token, value in mapped.items():
        if token in text:
            return value
    return ""


def guess_category(text: str) -> str:
    for category, tokens in CATEGORY_RULES:
        if any(token in text for token in tokens):
            return category
    return "기타"


def normalize_job_group(value: str) -> str:
    if "관리" in value:
        return "관리직"
    if "직접" in value:
        return "직접직"
    if "간접" in value:
        return "간접직"
    return value or ""



def normalize_iso_date(value: Any) -> str:
    text = norm(value)
    if not text:
        return ""
    explicit = re.search(r"^(20\d{2})[-./](\d{1,2})[-./](\d{1,2})", text)
    if explicit:
        year, month, day = explicit.groups()
        return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"
    short = re.search(r"^(\d{2})[-./](\d{1,2})[-./](\d{1,2})", text)
    if short:
        year, month, day = short.groups()
        return f"20{int(year):02d}-{int(month):02d}-{int(day):02d}"
    return extract_date(text)


def number_value(value: Any) -> float | None:
    text = norm(value).replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def positive_hours(value: Any) -> float | None:
    number = number_value(value)
    if number is None or number <= 0 or number > 24:
        return None
    return number


def positive_headcount(value: Any) -> int | None:
    number = number_value(value)
    if number is None or number <= 0 or number > 300:
        return None
    return int(round(number))


def structured_category(text: str) -> str:
    rules = (
        ("유지보수", ("보전", "설비", "점검", "경광등", "PLC", "칠러", "유압유", "누유", "도어락", "너트러너", "교체", "수정", "유지보수")),
        ("납품대응", ("납품", "출하", "상차", "불출", "물류", "포장")),
        ("납품용기/제품관리", ("선별", "검사", "불량", "반송품", "제품", "용기", "정기검사", "Burr", "BURR", "밀림", "조사")),
        ("설비/시설 청소", ("청소", "이형제")),
        ("행정업무", ("전산", "ISIR", "지적", "집계", "월마감", "서류", "자료", "보고", "재고조사")),
        ("개발/공정개선", ("개선", "개조", "C/T", "샘플", "이관설비", "T/O")),
        ("생산 돌대응", ("관리", "총괄", "지원판단", "품질문제", "공정관리", "현장관리", "대응")),
    )
    for category, tokens in rules:
        if any(token.lower() in text.lower() for token in tokens):
            return category
    return "기타"


def row_has_line_workbook_header(worksheet: Any) -> bool:
    header = row_to_text([worksheet.cell(2, col).value for col in range(1, min(worksheet.max_column, 18) + 1)])
    subheader = row_to_text([worksheet.cell(4, col).value for col in range(1, min(worksheet.max_column, 18) + 1)])
    return "NO" in header and "라인명" in header and "성" in header and "근" in header and "직접부서" in subheader and "간접부서" in subheader


def structured_candidate(
    source: dict[str, Any],
    source_id: str,
    sheet_name: str,
    date: str,
    row_number: int,
    team: str,
    name: str,
    job_group: str,
    headcount: int,
    hours: float | None,
    detail: str,
    category: str,
    parser: str,
) -> dict[str, Any]:
    return {
        "source_id": source_id,
        "date": date,
        "source_factory": extract_factory(f"{source['folder']} {source['file']} {sheet_name}", source["folder"]),
        "target_factory": "",
        "team": team,
        "name": name,
        "job_group": job_group,
        "headcount": str(headcount),
        "hours": "" if hours is None else f"{hours:g}",
        "detail": detail,
        "category_guess": category,
        "category_raw": category,
        "category_normalized": category,
        "confidence": "medium",
        "evidence": f"{source_id} R{row_number}: {team} / {name} / {job_group} / {headcount}명 / {'' if hours is None else f'{hours:g}h'} / {detail}",
        "needs_review_reason": f"{parser}_structured_candidate_confirm_before_report",
        "parser": parser,
    }


def declaration_candidates_from_sheet(source: dict[str, Any], source_id: str, worksheet: Any) -> list[dict[str, Any]]:
    text_top = " ".join(
        norm(worksheet.cell(row, col).value)
        for row in range(1, min(9, worksheet.max_row) + 1)
        for col in range(1, min(25, worksheet.max_column) + 1)
    )
    if "특근계획 신고서" not in text_top or "특근일자" not in text_top:
        return []
    date = ""
    factory = ""
    for col in range(1, min(25, worksheet.max_column) + 1):
        cell_text = norm(worksheet.cell(5, col).value)
        if cell_text == "공장 / 팀":
            factory = norm(worksheet.cell(5, col + 2).value)
        if cell_text == "특근일자":
            date = normalize_iso_date(worksheet.cell(5, col + 2).value)
    if not date:
        date = extract_date(text_top)
    if date not in TARGET_ISO_DATES:
        return []

    candidates: list[dict[str, Any]] = []
    current_team = ""
    current_job = ""
    current_work = ""
    current_detail = ""
    current_name = ""
    for row_number in range(9, worksheet.max_row + 1):
        first = norm(worksheet.cell(row_number, 1).value)
        if first.startswith("※") or first in {"합계", "TOTAL"}:
            break

        team = norm(worksheet.cell(row_number, 2).value) or current_team
        job_left = " ".join(
            part for part in (norm(worksheet.cell(row_number, 3).value), norm(worksheet.cell(row_number, 4).value)) if part
        )
        job_group = normalize_job_group(job_left or current_job)
        work = norm(worksheet.cell(row_number, 5).value) or current_work
        detail = norm(worksheet.cell(row_number, 9).value) or current_detail or work
        headcount = positive_headcount(worksheet.cell(row_number, 18).value) or positive_headcount(worksheet.cell(row_number, 14).value)
        hours = positive_hours(worksheet.cell(row_number, 19).value) or positive_hours(worksheet.cell(row_number, 15).value)
        note = norm(worksheet.cell(row_number, 22).value)
        name = ""
        if job_group == "관리직" and note:
            name = re.split(r"\s+", note)[0].strip()
            if name in {"주간", "야간", "관리직"}:
                name = ""

        if team:
            current_team = team
        if job_group:
            current_job = job_group
        if work:
            current_work = work
        if detail:
            current_detail = detail
        if name:
            current_name = name
        elif job_group == "관리직":
            name = current_name

        if not headcount or not detail or not team:
            continue
        if job_group != "관리직" and not any(token.lower() in f"{team} {work} {detail}".lower() for token in STRUCTURED_REPORTABLE_TEXT):
            continue
        if job_group == "직접직" and not any(token in detail for token in ("선별", "청소", "불량", "반송품", "검사")):
            continue

        category = structured_category(f"{team} {work} {detail}")
        candidates.append(
            structured_candidate(
                source,
                source_id,
                worksheet.title,
                date,
                row_number,
                team or factory,
                name,
                job_group,
                headcount,
                hours,
                detail,
                category,
                "declaration_sheet",
            )
        )
    return candidates


def management_candidates_from_sheet(source: dict[str, Any], source_id: str, worksheet: Any) -> list[dict[str, Any]]:
    if "관리직" not in worksheet.title:
        return []
    target_columns: list[tuple[int, str]] = []
    for col in range(1, min(worksheet.max_column, 80) + 1):
        date = normalize_iso_date(worksheet.cell(3, col).value)
        if date in TARGET_ISO_DATES:
            target_columns.append((col, date))
    if not target_columns:
        return []

    work_col = None
    for col in range(1, min(worksheet.max_column, 80) + 1):
        if "근무내용" in norm(worksheet.cell(3, col).value):
            work_col = col
            break
    candidates: list[dict[str, Any]] = []
    current_team = ""
    for row_number in range(5, worksheet.max_row + 1):
        team = norm(worksheet.cell(row_number, 1).value) or current_team
        if team and "합계" not in team:
            current_team = team
        position = norm(worksheet.cell(row_number, 2).value)
        name = norm(worksheet.cell(row_number, 3).value)
        if position:
            pass
        if not name or name in LINE_SUMMARY_NAMES or "합계" in team:
            continue
        detail = norm(worksheet.cell(row_number, work_col).value) if work_col else ""
        for col, date in target_columns:
            hours = positive_hours(worksheet.cell(row_number, col).value)
            if hours is None:
                continue
            candidate_detail = detail or f"{team} 관리직 특근"
            candidates.append(
                structured_candidate(
                    source,
                    source_id,
                    worksheet.title,
                    date,
                    row_number,
                    team,
                    name,
                    "관리직",
                    1,
                    hours,
                    candidate_detail,
                    structured_category(f"{team} {candidate_detail}"),
                    "management_sheet",
                )
            )
    return candidates


def line_workbook_candidates_from_sheet(source: dict[str, Any], source_id: str, worksheet: Any) -> list[dict[str, Any]]:
    if not row_has_line_workbook_header(worksheet):
        return []
    date = normalize_iso_date(worksheet.cell(1, 1).value) or extract_date(worksheet.title)
    if date not in TARGET_ISO_DATES:
        return []

    candidates: list[dict[str, Any]] = []
    current_direct_team = ""
    current_indirect_team = ""
    for row_number in range(5, worksheet.max_row + 1):
        direct_team = norm(worksheet.cell(row_number, 2).value) or current_direct_team
        if direct_team:
            current_direct_team = direct_team
        indirect_team = norm(worksheet.cell(row_number, 12).value) or current_indirect_team
        if indirect_team:
            current_indirect_team = indirect_team

        for block, team, name_col, detail_col, hours_col, job_group in (
            ("direct", direct_team, 3, 8, 9, "직접직"),
            ("indirect", indirect_team, 13, 16, 17, "간접직"),
        ):
            name = norm(worksheet.cell(row_number, name_col).value)
            detail = norm(worksheet.cell(row_number, detail_col).value)
            if not name or name in LINE_SUMMARY_NAMES or not detail or detail.isdigit():
                continue
            hours = positive_hours(worksheet.cell(row_number, hours_col).value)
            text = f"{team} {name} {detail}"
            if not any(token.lower() in text.lower() for token in STRUCTURED_REPORTABLE_TEXT):
                continue
            if block == "direct" and not any(token in detail for token in ("선별", "청소", "불량", "반송품", "검사", "개조", "교체", "점검")):
                continue
            if block == "indirect" and any(token in detail for token in ("정밀측정", "공정검사", "측정대응", "공구셋팅", "원통 연마", "드릴 호닝", "5축 연마")):
                continue
            candidates.append(
                structured_candidate(
                    source,
                    source_id,
                    worksheet.title,
                    date,
                    row_number,
                    team,
                    name,
                    job_group,
                    1,
                    hours,
                    detail,
                    structured_category(text),
                    "line_workbook_sheet",
                )
            )
    return candidates


def import_dependency_status() -> dict[str, bool]:
    status: dict[str, bool] = {}
    for name in ("openpyxl", "pypdf"):
        try:
            __import__(name)
            status[name] = True
        except Exception:
            status[name] = False
    return status


def file_inventory(reply_dir: Path) -> list[dict[str, Any]]:
    files = sorted([path for path in reply_dir.rglob("*") if path.is_file()], key=lambda p: (str(p.parent), p.name))
    items = []
    for path in files:
        stat = path.stat()
        items.append(
            {
                "folder": path.parent.name,
                "file": path.name,
                "path": str(path),
                "relative_path": rel(path, reply_dir.parent),
                "extension": path.suffix.lower(),
                "size_bytes": stat.st_size,
                "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
                "role_guess": role_guess(path),
            }
        )
    return items


def build_evidence(source: dict[str, Any], source_id: str, sheet_or_page: str, location: str, text: str, text_kind: str) -> dict[str, Any]:
    keyword_hits = hits(text, KEYWORDS)
    date_hits = hits(text, TARGET_DATE_TOKENS)
    return {
        "source_id": source_id,
        "source_file": source["file"],
        "source_path": source["relative_path"],
        "source_folder": source["folder"],
        "source_sheet_or_page": sheet_or_page,
        "source_location": location,
        "source_type": text_kind,
        "text": text[:12000],
        "text_chars": len(text),
        "keyword_hits": keyword_hits,
        "date_hits": date_hits,
        "current_period": bool(date_hits or contains_any(sheet_or_page, TARGET_DATE_TOKENS) or contains_any(source["file"], TARGET_DATE_TOKENS)),
        "factory_guess": extract_factory(f"{source['folder']} {source['file']} {sheet_or_page} {text}", source["folder"]),
    }


def extract_xlsx(source: dict[str, Any], source_index: int, reply_dir: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    try:
        import openpyxl
    except Exception as exc:
        return [], {"status": "dependency_missing", "error": f"openpyxl missing: {exc}"}

    path = Path(source["path"])
    evidence: list[dict[str, Any]] = []
    sheet_summaries: list[dict[str, Any]] = []
    local_candidates: list[dict[str, Any]] = []
    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
            evidence_lines: list[str] = []
            keyword_counter: Counter[str] = Counter()
            date_counter: Counter[str] = Counter()
            rows_scanned = 0
            non_empty = 0
            cached_rows: list[list[Any]] = []
            for row_number, row in enumerate(worksheet.iter_rows(max_row=220, max_col=60, values_only=True), start=1):
                rows_scanned += 1
                row_values = list(row)
                cached_rows.append(row_values)
                text = row_to_text(row_values)
                if not text:
                    continue
                non_empty += 1
                row_keywords = hits(text, KEYWORDS)
                row_dates = hits(text, TARGET_DATE_TOKENS)
                keyword_counter.update(row_keywords)
                date_counter.update(row_dates)
                if len(evidence_lines) < 42 and (row_keywords or row_dates or len(evidence_lines) < 4):
                    evidence_lines.append(f"R{row_number}: {text}")
            sheet_summary = {
                "sheet": worksheet.title,
                "rows_scanned": rows_scanned,
                "non_empty_rows": non_empty,
                "keyword_hits": dict(keyword_counter.most_common(20)),
                "date_hits": dict(date_counter.most_common(20)),
                "current_period": bool(date_counter or contains_any(worksheet.title, TARGET_DATE_TOKENS)),
                "evidence_line_count": len(evidence_lines),
            }
            sheet_summaries.append(sheet_summary)
            if evidence_lines and (sheet_summary["current_period"] or keyword_counter):
                source_id = f"xlsx-{source_index:03d}-{sheet_index:03d}"
                text = "\n".join(
                    [
                        f"file={source['file']}",
                        f"sheet={worksheet.title}",
                        f"folder={source['folder']}",
                        *evidence_lines,
                    ]
                )
                evidence.append(build_evidence(source, source_id, worksheet.title, f"sheet:{worksheet.title}", text, "xlsx_sheet"))
                sheet_snapshot = WorksheetSnapshot(worksheet.title, cached_rows, min(worksheet.max_column, 60))
                if "특근계획 신고서" in text:
                    local_candidates.extend(declaration_candidates_from_sheet(source, source_id, sheet_snapshot))
                if "관리직" in worksheet.title:
                    local_candidates.extend(management_candidates_from_sheet(source, source_id, sheet_snapshot))
                if "직접부서" in text and "간접부서" in text:
                    local_candidates.extend(line_workbook_candidates_from_sheet(source, source_id, sheet_snapshot))
    except Exception as exc:
        return evidence, {"status": "error", "error": repr(exc), "sheets": sheet_summaries, "local_candidate_rows": local_candidates}
    return evidence, {"status": "ok", "sheets": sheet_summaries, "sheet_count": len(sheet_summaries), "local_candidate_rows": local_candidates, "local_candidate_count": len(local_candidates)}


def extract_pdf(source: dict[str, Any], source_index: int) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        return [], {"status": "dependency_missing", "error": f"pypdf missing: {exc}"}

    path = Path(source["path"])
    evidence: list[dict[str, Any]] = []
    pages: list[dict[str, Any]] = []
    try:
        reader = PdfReader(str(path))
        for page_index, page in enumerate(reader.pages, start=1):
            text = " ".join((page.extract_text() or "").split())
            page_summary = {
                "page": page_index,
                "text_chars": len(text),
                "keyword_hits": hits(text, KEYWORDS),
                "date_hits": hits(text, TARGET_DATE_TOKENS),
                "current_period": contains_any(text, TARGET_DATE_TOKENS),
            }
            pages.append(page_summary)
            if text and (page_summary["keyword_hits"] or page_summary["date_hits"]):
                source_id = f"pdf-{source_index:03d}-{page_index:03d}"
                evidence.append(build_evidence(source, source_id, f"p{page_index}", f"page:{page_index}", text, "pdf_page"))
    except Exception as exc:
        return evidence, {"status": "error", "error": repr(exc), "pages": pages}
    return evidence, {"status": "ok", "pages": pages, "page_count": len(pages)}


def extract_pptx(source: dict[str, Any], source_index: int) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    path = Path(source["path"])
    evidence: list[dict[str, Any]] = []
    slides: list[dict[str, Any]] = []
    try:
        with zipfile.ZipFile(path) as archive:
            names = sorted(
                [name for name in archive.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", name)],
                key=lambda name: int(re.search(r"slide(\d+)\.xml", name).group(1)),
            )
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            for slide_index, name in enumerate(names, start=1):
                root = ET.fromstring(archive.read(name))
                text_parts = [node.text.strip() for node in root.findall(".//a:t", ns) if node.text and node.text.strip()]
                text = " ".join(text_parts)
                slide_summary = {
                    "slide": slide_index,
                    "text_chars": len(text),
                    "keyword_hits": hits(text, KEYWORDS),
                    "date_hits": hits(text, TARGET_DATE_TOKENS),
                    "current_period": contains_any(text, TARGET_DATE_TOKENS),
                }
                slides.append(slide_summary)
                if text and (slide_summary["keyword_hits"] or slide_summary["date_hits"]):
                    source_id = f"pptx-{source_index:03d}-{slide_index:03d}"
                    evidence.append(build_evidence(source, source_id, f"slide {slide_index}", f"slide:{slide_index}", text, "pptx_slide"))
    except Exception as exc:
        return evidence, {"status": "error", "error": repr(exc), "slides": slides}
    return evidence, {"status": "ok", "slides": slides, "slide_count": len(slides)}


LABELED_ROW_PATTERN = re.compile(
    r"(?:부서명\s*)?(?P<team>[^/\n]{1,30})\s*/\s*(?:직군\s*)?(?P<job>관리직|간접직|직접직|관리|간접|직접)\s*/.*?(?:업무명\s*)?(?P<work>[^/\n]{1,60})\s*/\s*(?:상세\s*내역\s*)?(?P<detail>[^/\n]{1,120})\s*/\s*(?:인원\s*)?(?P<headcount>\d+)\s*/\s*(?:시간\s*)?(?P<hours>\d+(?:\.\d+)?)",
    flags=re.IGNORECASE,
)


def local_candidates_from_evidence(evidence_items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str, str, str]] = set()
    for item in evidence_items:
        if not item.get("current_period"):
            continue
        text = item["text"]
        date = extract_date(text)
        if date and date not in TARGET_ISO_DATES:
            continue
        source_factory = extract_factory(text, item.get("factory_guess", ""))
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            match = LABELED_ROW_PATTERN.search(line)
            if not match:
                continue
            team = re.sub(r"^R\d+:\s*", "", match.group("team")).strip()
            job_group = normalize_job_group(match.group("job"))
            detail = match.group("detail").strip()
            headcount = match.group("headcount")
            hours = match.group("hours")
            try:
                headcount_number = float(headcount)
                hours_number = float(hours)
            except ValueError:
                continue
            if not (0 < headcount_number <= 200 and 0 < hours_number <= 24):
                continue
            if team.isnumeric() or detail.isnumeric() or team in {"NO.", "NO", "직군", "현장직"}:
                continue
            if not any(token in line for token in ("부서", "직군", "업무", "상세", "관리직", "간접직", "직접직")):
                continue
            key = (item["source_id"], date, team, headcount, detail)
            if key in seen:
                continue
            seen.add(key)
            candidates.append(
                {
                    "source_id": item["source_id"],
                    "date": date,
                    "source_factory": source_factory,
                    "target_factory": "",
                    "team": team,
                    "name": "",
                    "job_group": job_group,
                    "headcount": headcount,
                    "hours": hours,
                    "detail": detail,
                    "category_guess": guess_category(f"{team} {match.group('work')} {detail}"),
                    "confidence": "medium",
                    "evidence": line[:500],
                    "needs_review_reason": "local_regex_candidate_confirm_before_report",
                }
            )
    return candidates


def dedupe_local_candidates(candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    deduped: list[dict[str, Any]] = []
    seen: set[tuple[str, ...]] = set()
    parser_rank = {
        "declaration_sheet": 0,
        "management_sheet": 1,
        "line_workbook_sheet": 2,
    }
    sorted_candidates = sorted(candidates, key=lambda row: parser_rank.get(str(row.get("parser") or ""), 9))
    for row in sorted_candidates:
        date = norm(row.get("date"))
        name = norm(row.get("name"))
        detail = norm(row.get("detail"))
        team = norm(row.get("team"))
        source_id = norm(row.get("source_id"))
        if name:
            key = (date, name)
        else:
            key = (date, source_id, team, detail, norm(row.get("headcount")))
        if key in seen:
            continue
        seen.add(key)
        deduped.append(row)
    return deduped


def run_reply_import(reply_dir: Path, period_label: str = "2026년 5월 5주차") -> dict[str, Any]:
    run_id = f"imp-{now_id()}"
    reply_dir = reply_dir.resolve()
    started_at = datetime.now().isoformat(timespec="seconds")
    inventory = file_inventory(reply_dir)
    evidence_items: list[dict[str, Any]] = []
    sources: list[dict[str, Any]] = []
    structured_local_candidates: list[dict[str, Any]] = []

    for source_index, source in enumerate(inventory, start=1):
        ext = source["extension"]
        extracted: list[dict[str, Any]] = []
        details: dict[str, Any]
        if ext == ".xlsx":
            extracted, details = extract_xlsx(source, source_index, reply_dir)
        elif ext == ".pdf":
            extracted, details = extract_pdf(source, source_index)
        elif ext == ".pptx":
            extracted, details = extract_pptx(source, source_index)
        elif ext == ".xls":
            details = {
                "status": "unsupported",
                "error": "legacy .xls requires resave to .xlsx or controlled Excel COM conversion",
            }
        else:
            details = {"status": "skipped", "error": f"unsupported extension {ext}"}
        evidence_items.extend(extracted)
        structured_local_candidates.extend(details.get("local_candidate_rows", []))
        sources.append(
            {
                **source,
                "extract_status": details.get("status", "unknown"),
                "extract_error": details.get("error", ""),
                "evidence_count": len(extracted),
                "details": details,
            }
        )

    extension_counts = Counter(item["extension"] for item in inventory)
    folder_counts = Counter(item["folder"] for item in inventory)
    status_counts = Counter(item["extract_status"] for item in sources)
    evidence_type_counts = Counter(item["source_type"] for item in evidence_items)
    current_evidence_count = sum(1 for item in evidence_items if item["current_period"])
    local_candidates = dedupe_local_candidates([*local_candidates_from_evidence(evidence_items), *structured_local_candidates])
    finished_at = datetime.now().isoformat(timespec="seconds")
    return {
        "run_id": run_id,
        "period_label": period_label,
        "reply_dir": str(reply_dir),
        "started_at": started_at,
        "finished_at": finished_at,
        "mode": "local_extract_only",
        "dependency_status": import_dependency_status(),
        "summary": {
            "file_count": len(inventory),
            "extension_counts": dict(extension_counts),
            "folder_counts": dict(folder_counts),
            "extract_status_counts": dict(status_counts),
            "evidence_count": len(evidence_items),
            "current_period_evidence_count": current_evidence_count,
            "evidence_type_counts": dict(evidence_type_counts),
            "local_candidate_count": len(local_candidates),
            "unsupported_count": status_counts.get("unsupported", 0),
        },
        "sources": sources,
        "evidence_items": evidence_items,
        "local_candidate_rows": local_candidates,
        "ai": {
            "enabled": False,
            "candidate_rows": [],
            "source_assessment": [],
            "conflicts": [],
            "questions_for_human": [],
            "abstentions": [],
            "errors": [],
            "usage_metadata": [],
        },
        "review_policy": "candidate_rows_are_not_final_report_data_until_human_review_lock",
    }


def save_run(run: dict[str, Any], output_dir: Path) -> dict[str, str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    run_path = output_dir / f"{run['run_id']}.json"
    latest_path = output_dir / "latest.json"
    text = json.dumps(run, ensure_ascii=False, indent=2)
    run_path.write_text(text, encoding="utf-8")
    latest_path.write_text(text, encoding="utf-8")
    return {"run_path": str(run_path), "latest_path": str(latest_path)}


def load_latest(output_dir: Path) -> dict[str, Any] | None:
    latest_path = output_dir / "latest.json"
    if not latest_path.exists():
        return None
    return json.loads(latest_path.read_text(encoding="utf-8"))
